#!/usr/bin/env python3
"""Build berlin-wall-sources.docx from the same verified source data as the PDF."""

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

RED = RGBColor(0xA8, 0x23, 0x1F)
GREY = RGBColor(0x66, 0x66, 0x66)
INK = RGBColor(0x1A, 0x1A, 0x1A)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def run(p, text, *, bold=False, italic=False, size=10, color=None, font="Calibri"):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.size = Pt(size)
    r.font.name = font
    if color is not None:
        r.font.color.rgb = color
    return r


def rich(p, parts, size=10, font="Calibri"):
    """parts: list of (text, style) where style in {'', 'b', 'i', 'bi'}"""
    for text, style in parts:
        run(p, text, bold="b" in style, italic="i" in style, size=size, font=font)


def label(cell, text):
    p = cell.paragraphs[0] if not cell.paragraphs[0].text else cell.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run(p, text, bold=True, size=7.5, color=GREY)
    return p


doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(1.6)
    s.left_margin = s.right_margin = Cm(1.6)

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10)

# ---------------- title ----------------
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
run(p, "The Berlin Wall, 1961–1989", bold=True, size=26, color=RED)

p = doc.add_paragraph()
rich(p, [("Source evidence pack. ", "b"),
         ("Ten real primary and secondary sources, the evidence quoted from each, "
          "what it proves, how reliable it is, and where it came from.", "")], size=11)

p = doc.add_paragraph()
rich(p, [("How to use this document. ", "b"),
         ("Print it and cut out the tables, or edit them in Word to fit your own chart. "
          "The grey box in each source is where a photograph goes — open ", ""),
         ("photo-sheet.html", "b"),
         (" in a web browser to get the eight photographs (P1–P8) with their credits "
          "and copyright licences.", "")])

p = doc.add_paragraph()
run(p, "About the photographs: they are not embedded in this document because they are stored "
       "on archive websites. All eight are genuine and free to use for schoolwork — the CIA "
       "photographs are public domain, the Bundesarchiv photographs are CC BY-SA and must credit "
       "the Bundesarchiv.", size=8.5, color=GREY)


def heading(text):
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run(p, text, bold=True, size=15, color=INK)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:color"), "1A1A1A")
    pbdr.append(bottom)
    pPr.append(pbdr)


# ---------------- definitions ----------------
heading("Definitions")

DEFS = [
    ("Primary source", "Evidence created at the time by someone who was there. It has not been filtered through a later writer.", "S1, S2, S3, S4, S5, S6, S8, S9"),
    ("Secondary source", "Evidence made later by someone studying the event, usually by gathering and checking many primary sources.", "S7 and S10 — both compiled by historians after 1990"),
    ("Provenance", "Where a source came from: who made it, when, for whom, and why. Judge a source by its provenance before you judge its content.", "S5 was written by the officer whose own men did the shooting"),
    ("Bias", "When a source leans one way because of who made it. Bias does not make a source useless — it makes it evidence of what that side wanted people to believe.", "S4 is pure propaganda, and that is why it is good evidence"),
    ("Reliability", "How far you can trust that what a source says actually happened. A source can be unreliable about facts but reliable about attitudes.", "S3 is unreliable as fact, reliable as the official line"),
    ("Corroboration", "Checking one source against another. Two sources from opposite sides that agree on a fact make that fact very strong.", "S1 and S2 are both East German — and they contradict each other"),
    ("Utility", "How useful a source is for the specific question you are answering. A source can be biased and still be highly useful.", "Always name the question first"),
    ("Propaganda", "Information published deliberately to make people support a government or cause, whether or not it is true.", "S3 and S4 are both official East German propaganda"),
    ("GDR / DDR", "The German Democratic Republic (East Germany), 1949–1990. A communist one-party state allied to the Soviet Union. It built the Wall.", "Ruled by the SED, the Socialist Unity Party"),
    ("FRG / BRD", "The Federal Republic of Germany (West Germany), democratic and capitalist, allied to the USA, Britain and France.", "West Berlin sat about 160 km inside East Germany"),
    ("Republikflucht", "“Flight from the republic” — the East German legal term for leaving without permission. It was a crime.", "About 2.7–3.5 million people left before 1961"),
    ("Antifaschistischer Schutzwall", "“Anti-fascist protection rampart” — the official East German name for the Wall. It claimed the Wall kept fascists out, not citizens in.", "Ordinary East Germans never used the phrase"),
    ("The death strip", "The cleared ground behind the outer wall: raked sand, floodlights, dog runs, trip-wire fences, anti-vehicle trenches and watchtowers.", "The system was roughly 30–100 m deep"),
    ("Checkpoint Charlie", "The crossing point on Friedrichstraße used by Allied staff and foreigners. In October 1961 US and Soviet tanks faced each other here.", "Peter Fechter was shot about 200 m away (S5)"),
]

t = doc.add_table(rows=1, cols=3)
t.style = "Table Grid"
t.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = t.rows[0].cells
for i, h in enumerate(["Term", "What it means", "Example in this pack"]):
    hdr[i].text = ""
    run(hdr[i].paragraphs[0], h, bold=True, size=9, color=RGBColor(0xFF, 0xFF, 0xFF))
    shade(hdr[i], "1A1A1A")
for term, meaning, eg in DEFS:
    c = t.add_row().cells
    c[0].text = ""
    run(c[0].paragraphs[0], term, bold=True, size=9.5)
    c[1].text = ""
    run(c[1].paragraphs[0], meaning, size=9.5)
    c[2].text = ""
    run(c[2].paragraphs[0], eg, size=9)

# ---------------- source cards ----------------
heading("The ten sources")

SOURCES = [
    dict(
        sid="S1", type="PRESS CONFERENCE", year="15 JUNE 1961", side="EAST GERMANY",
        title="“Nobody intends to put up a wall”",
        origin=[("Walter Ulbricht", "b"), (", leader of East Germany, answering a West German journalist at an "
                "international press conference in the House of Ministries, East Berlin. Printed in the party newspaper ", ""),
                ("Neues Deutschland", "i"), (" the next day.", "")],
        quote="“Ich verstehe Ihre Frage so, daß es in Westdeutschland Menschen gibt, die wünschen, "
              "daß wir die Bauarbeiter der DDR dazu mobilisieren eine Mauer aufzurichten. Mir ist nicht "
              "bekannt, daß eine solche Absicht besteht. Niemand hat die Absicht, eine Mauer zu errichten.”",
        trans="“I understand your question to mean that there are people in West Germany who wish that we would "
              "mobilise the construction workers of the GDR to put up a wall. I am not aware that any such intention "
              "exists. Nobody intends to put up a wall.”",
        shows="The East German leader publicly denied any plan for a wall 59 days before he built one. It proves the "
              "decision was kept secret, and it is the most famous lie of the Cold War.",
        rely="Completely reliable as a record of what was said — it was public and recorded. Completely unreliable "
             "as a statement of fact. Note that Ulbricht was the first to say the word “wall”.",
        cite="Walter Ulbricht, international press conference, East Berlin, 15 June 1961; transcript in Neues Deutschland, "
             "16 June 1961. Via German History in Documents and Images (GHDI).",
        photo="PHOTO P1", photocap="Conrad Schumann's leap over the barbed wire, 15 Aug 1961. CIA collection, public domain.",
    ),
    dict(
        sid="S2", type="SECRET GOVERNMENT RECORD", year="1 AUGUST 1961", side="USSR & EAST GERMANY",
        title="Khrushchev and Ulbricht plan the border closure",
        origin=[("Notes of a private conversation between ", ""), ("Nikita Khrushchev", "b"), (" (Soviet leader) and ", ""),
                ("Walter Ulbricht", "b"), (", taken twelve days before the Wall went up. Secret at the time; declassified "
                "after 1990 and published by the Wilson Center Digital Archive.", "")],
        quote="Khrushchev, on East Germany's problem: “Many engineers have fled the GDR. You should consider — "
              "shouldn't we send you perhaps some engineers from the Soviet Union? They won't run away. […] "
              "But something must be done.”",
        trans="Ulbricht presses for economic help from the other communist states in case the West retaliates once the "
              "border is shut.",
        shows="Behind closed doors the real reason was economic collapse through emigration — skilled workers leaving "
              "— not fascism or Western attack. It also shows Moscow, not just East Berlin, was in on the decision.",
        rely="Very high. It was never meant to be read by the public, so there was no reason to spin it. This is the "
             "source that flatly contradicts S1, S3 and S4.",
        cite="“Notes on the Conversation of Comrade N. S. Khrushchev with Comrade W. Ulbricht on 1 August 1961”, "
             "Wilson Center Digital Archive.",
        photo="PHOTO P2", photocap="Workers building the Wall under armed guard, 1961. CIA collection, public domain.",
    ),
    dict(
        sid="S3", type="NEWSPAPER ARTICLE", year="14 AUGUST 1961", side="EAST GERMANY",
        title="“Measures to protect peace” — the front page",
        origin=[("Front page of ", ""), ("Neues Deutschland", "bi"), (", the official daily newspaper of the SED (the East "
                "German ruling party), published the morning after the border was sealed. A state-controlled newspaper, "
                "not a free press.", "")],
        quote="“Zur Unterbindung der feindlichen Tätigkeit der revanchistischen und militaristischen Kräfte "
              "Westdeutschlands und Westberlins wird eine solche Kontrolle an den Grenzen der Deutschen Demokratischen "
              "Republik … eingeführt, wie sie an den Grenzen jedes souveränen Staates üblich ist.”",
        trans="“To stop the hostile activity of the revanchist and militarist forces of West Germany and West Berlin, "
              "controls are being introduced at the borders of the German Democratic Republic … such as are normal at "
              "the borders of any sovereign state.”",
        shows="The official public justification: the Wall was defensive and normal. The paper also claims most Berliners "
              "approved — useful evidence of how the state managed the news.",
        rely="Low as fact — every East German newspaper was censored by the party. High as evidence of the official "
             "line. Compare it with S2, written thirteen days earlier in private.",
        cite="“Maßnahmen zum Schutz des Friedens und zur Sicherung der Deutschen Demokratischen Republik in "
             "Kraft”, Neues Deutschland, 14 August 1961, p. 1.",
        photo="PHOTO P3", photocap="People in Leipzig read a newspaper special edition about the Wall, Aug 1961. "
                                   "Bundesarchiv, CC BY-SA.",
    ),
    dict(
        sid="S4", type="PROPAGANDA PAMPHLET", year="1961", side="EAST GERMANY",
        title="“In praise of the Berlin Wall”",
        origin=[("A 48-page illustrated pamphlet issued by the ", ""), ("SED", "b"), (" within weeks of 13 August 1961, "
                "produced by the “Fighting Group Berlin-Mitte” and Berlin publishing houses. Handed out to "
                "workers and party members.", "")],
        quote="“On 13 August 1961 peace-loving Berliners won a battle for peace. The battle groups of Berlin's working "
              "class … put an end to subversive activity against the GDR by spies, slave-traders, and Revanchist "
              "organizations based in West Berlin.”",
        trans="The back cover reads: “West German militarism has been frustrated by the building of the Berlin Wall.”",
        shows="How the state sold the Wall to its own people: as a victory, not a restriction. The language — "
              "“spies”, “slave-traders” — shows the enemy being invented to justify the barrier.",
        rely="Worthless as a factual account. Excellent as evidence of propaganda technique. Notice it never mentions the "
             "2.7 million people who had already left (see S7).",
        cite="SED pamphlet, 1961, translated as “In Praise of the Berlin Wall”, German Propaganda Archive, "
             "Calvin University (ed. Randall Bytwerk).",
        photo="PHOTO P4", photocap="The Wall at the Brandenburg Gate, 1961. Bundesarchiv, CC BY-SA.",
    ),
    dict(
        sid="S5", type="MILITARY REPORT (SECRET)", year="17 AUGUST 1962", side="EAST GERMANY",
        title="The border commander's report on the shooting of Peter Fechter",
        origin=[("Internal report by the commander of the ", ""), ("1st Border Brigade", "b"), (" on the day 18-year-old "
                "bricklayer ", ""), ("Peter Fechter", "b"), (" was shot climbing the Wall near Checkpoint Charlie. "
                "Classified at the time; now held at the German Federal Archive.", "")],
        quote="“Die zweite Person wurde beim Anspringen an die Mauer getroffen … die Person brach daraufhin an "
              "der Mauer zusammen und schrie um Hilfe.”",
        trans="“The second person was hit as he jumped up at the Wall … the person then collapsed at the Wall and "
              "screamed for help.” Four guards fired 35 shots. Fechter bled to death over roughly an hour while both "
              "sides refused to cross. A smoke screen was laid to recover him. He was declared dead at 17:00.",
        shows="The shoot-to-kill order in practice, described by the army itself. The report's conclusions are about "
              "plugging the gap and bricking up windows — not about the death of a teenager.",
        rely="Written by the side that did the shooting, to protect itself — it claims warning shots were fired, which "
             "the later trial disputed. But it is a secret internal document, so the basic facts are strong.",
        cite="“Bericht des Kommandeurs der 1. Grenzbrigade über den Fluchtversuch und die Erschießung von "
             "Peter Fechter”, 17 August 1962. BArch, VA-07/16930, Bl. 2–5. Via Chronik der Mauer.",
        photo="PHOTO P5", photocap="Armed East German border guards. CIA collection, public domain. (The photo of Fechter "
                                   "himself is still in copyright.)",
    ),
    dict(
        sid="S6", type="POLITICAL SPEECH", year="26 JUNE 1963", side="USA / WEST",
        title="Kennedy: “Ich bin ein Berliner”",
        origin=[("Speech by US President ", ""), ("John F. Kennedy", "b"), (" to a crowd of several hundred thousand at the "
                "Rudolph Wilde Platz, West Berlin. Filmed, broadcast and transcribed; the record is held by the US National "
                "Archives.", "")],
        quote="“Freedom has many difficulties and democracy is not perfect, but we have never had to put a wall up to "
              "keep our people in, to prevent them from leaving us. … While the wall is the most obvious and vivid "
              "demonstration of the failures of the Communist system, for all the world to see, we take no satisfaction in "
              "it, for it is … an offense not only against history but an offense against humanity.”",
        trans="",
        shows="The Western argument in one line: a state that has to lock its people in has already lost the argument. It "
              "also shows the Wall being used as a propaganda weapon by both sides.",
        rely="Reliable as a record of what was said. But it is a political speech to a friendly crowd — designed to "
             "reassure West Berliners, not to analyse. Kennedy did not act to remove the Wall.",
        cite="John F. Kennedy, Remarks at the Rudolph Wilde Platz, Berlin, 26 June 1963. US National Archives; transcript "
             "via the Miller Center and the JFK Presidential Library.",
        photo="PHOTO P6", photocap="A family separated by the Wall — children held up to be seen from the other side. "
                                   "CIA collection, public domain.",
    ),
    dict(
        sid="S7", type="STATISTICS", year="1949–1990", side="SECONDARY / RESEARCH",
        title="Migration figures: East to West Germany",
        origin=[("Official migration statistics compiled from ", ""), ("West German and East German government records", "b"),
                (", published as a data table by German History in Documents and Images (GHDI), a project of the German "
                 "Historical Institute.", "")],
        quote="People moving to West Germany, by year — 1953: 331,390 · 1960: 199,188 · 1961: 207,026 · "
              "1962: 21,365 · 1963: 42,622 · 1988: 39,832 · 1989: 343,854",
        trans="Around 2.7 million people left between 1949 and August 1961. After the Wall went up, the number fell by "
              "roughly 90% in a single year.",
        shows="Hard proof that the Wall worked, and proof of why it was built. The 1989 spike shows the system failing "
              "again the moment the border reopened.",
        rely="High — two governments counted independently and the totals broadly agree. But numbers alone hide the "
             "human reason people left; pair this with S5 and S10.",
        cite="“East-West German Immigration Statistics (1949–1990)” and “Refugee Movement "
             "(1950–1963)”, German History in Documents and Images, German Historical Institute.",
        photo="GRAPH", photocap="Draw a bar chart from the figures on the left. The drop from 1961 to 1962 is the point "
                                "of the graph.",
    ),
    dict(
        sid="S8", type="POLITICAL SPEECH", year="12 JUNE 1987", side="USA / WEST",
        title="Reagan: “Mr. Gorbachev, tear down this wall!”",
        origin=[("Speech by US President ", ""), ("Ronald Reagan", "b"), (" at the Brandenburg Gate, West Berlin, delivered "
                "behind bulletproof glass and amplified so it could be heard on the eastern side. Official text held by the "
                "Reagan Presidential Library.", "")],
        quote="“General Secretary Gorbachev, if you seek peace, if you seek prosperity for the Soviet Union and Eastern "
              "Europe, if you seek liberalization: Come here to this gate! Mr. Gorbachev, open this gate! Mr. Gorbachev, "
              "tear down this wall!”",
        trans="Earlier in the same speech: “Every man is a Berliner, forced to look upon a scar.”",
        shows="The Wall still standing 26 years on, and the West using it as a public challenge. Reagan's own State "
              "Department tried to cut the line — evidence of a split in Western policy.",
        rely="Reliable text, but treat the popular claim that this speech caused the Wall to fall with caution. The Wall "
             "fell 29 months later for reasons largely inside the Eastern bloc (see S9).",
        cite="Ronald Reagan, “Remarks on East-West Relations at the Brandenburg Gate in West Berlin”, 12 June "
             "1987. Ronald Reagan Presidential Library & US National Archives.",
        photo="PHOTO P7", photocap="Checkpoint Charlie from the air. CIA collection, public domain.",
    ),
    dict(
        sid="S9", type="GOVERNMENT DECREE", year="9 NOVEMBER 1989", side="EAST GERMANY",
        title="“Schabowski's note” — the order that opened the Wall",
        origin=[("The draft decision of the ", ""), ("GDR Council of Ministers", "b"), (" on travel and permanent exit, "
                "circulated on 9 November 1989 and marked confidential. These were the notes Günter Schabowski read "
                "out, unprepared, at a live press conference that evening.", "")],
        quote="“Privatreisen nach dem Ausland können ohne Vorliegen von Voraussetzungen (Reiseanlässe und "
              "Verwandtschaftsverhältnisse) beantragt werden. … Ständige Ausreisen können über alle "
              "Grenzübergangsstellen der DDR zur BRD bzw. zu Berlin (West) erfolgen.”",
        trans="“Private trips abroad may be applied for without preconditions … Permanent exit may take place via "
              "all GDR border crossings to the FRG or to Berlin (West).” The document says the press release was to be "
              "published the next day, 10 November.",
        shows="The Wall opened partly by accident. The decree was meant to control emigration, not end it — and to "
              "start a day later. Asked when it applied, Schabowski said immediately, and crowds went to the crossings.",
        rely="An official internal document — very reliable for what was decided. It cannot tell you what Schabowski "
             "actually said on air; for that you need the press conference recording.",
        cite="“Zeitweilige Übergangsregelung für Reisen und ständige Ausreise aus der DDR”, "
             "9 November 1989. BStU, MfS, Arbeitsbereich Neiber Nr. 553, Bl. 15–19. Via Chronik der Mauer.",
        photo="PHOTO P8", photocap="Anti-tank defences inside the border strip. CIA collection, public domain.",
    ),
    dict(
        sid="S10", type="RESEARCH DATABASE", year="PUBLISHED 2017–19", side="SECONDARY / ACADEMIC",
        title="Victims at the Wall: the official death toll",
        origin=[("A joint research project by the ", ""), ("Berlin Wall Foundation", "b"), (" and the ", ""),
                ("Centre for Contemporary History (ZZF) Potsdam", "b"), (", funded by the German federal government. "
                "Researchers went through border-troop and Stasi files, 1990s trial records, and interviewed families.", "")],
        quote="“Between 1961 and 1989, at least 140 people were killed at the Berlin Wall or died under circumstances "
              "directly connected with the GDR border regime. In addition, at least 251 people from East and West died "
              "before, during or after controls at Berlin border crossings while travelling.”",
        trans="Altogether 576 cases were recorded and checked. Earlier counts had ranged from 78 to over 200, depending on "
              "the criteria used.",
        shows="The human cost, established by evidence rather than by either side's propaganda. It also shows how history "
              "gets settled: by defining terms and checking records.",
        rely="The most reliable source in this pack. But note it is secondary and it says “at least” — it "
             "excludes people who died of grief, and figures depend on the definition chosen.",
        cite="Hans-Hermann Hertle & Maria Nooke (eds.), The Victims at the Berlin Wall, 1961–1989: A Biographical "
             "Handbook, ZZF Potsdam & Berlin Wall Foundation; database at Chronik der Mauer.",
        photo="BREAKDOWN", photocap="101 killed while escaping · 30 not escaping · 8 border soldiers · "
                                    "1 Soviet soldier = 140 total",
    ),
]

for s in SOURCES:
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    t.autofit = False

    # header row (merged)
    hr = t.add_row().cells
    hc = hr[0].merge(hr[1])
    hc.text = ""
    p = hc.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run(p, f"{s['sid']}  |  {s['type']}  |  {s['year']}  |  {s['side']}", bold=True, size=8.5, color=RED)
    p2 = hc.add_paragraph()
    run(p2, s["title"], bold=True, size=12.5)
    shade(hc, "F0EDE8")

    # body rows; photo cell spans all four
    r = t.add_row().cells
    label(r[0], "ORIGIN")
    p = r[0].add_paragraph()
    rich(p, s["origin"], size=9.5)
    photo_cell = r[1]
    photo_cell.text = ""
    p = photo_cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, s["photo"], bold=True, size=11, color=RGBColor(0x3F, 0x4A, 0x55))
    if s["photo"].startswith("PHOTO"):
        p = photo_cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(p, "Glue here", size=9, color=GREY)
    p = photo_cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, s["photocap"], size=8, color=GREY)
    shade(photo_cell, "F4F3F1")

    r = t.add_row().cells
    label(r[0], "THE EVIDENCE")
    p = r[0].add_paragraph()
    run(p, s["quote"], size=10, font="Georgia")
    if s["trans"]:
        p = r[0].add_paragraph()
        run(p, s["trans"], italic=True, size=9, color=GREY)
    shade(r[0], "FDF7F6")

    r = t.add_row().cells
    label(r[0], "WHAT IT SHOWS")
    p = r[0].add_paragraph()
    run(p, s["shows"], size=9.5)
    label(r[0], "RELIABILITY")
    p = r[0].add_paragraph()
    run(p, s["rely"], size=9.5)

    r = t.add_row().cells
    p = r[0].paragraphs[0]
    run(p, "Cite as: ", bold=True, size=8, color=GREY)
    run(p, s["cite"], size=8, color=GREY)

    # merge the photo column down across the four body rows
    rows = t.rows
    photo_col = [rows[i].cells[1] for i in range(1, len(rows))]
    merged = photo_col[0]
    for c in photo_col[1:]:
        merged = merged.merge(c)

    for row in t.rows:
        row.cells[0].width = Cm(12.6)
    doc.add_paragraph()

# ---------------- master chart ----------------
heading("Master chart")
CHART = [
    ("S1", "Press conference", "15 Jun 1961", "East Germany", "The leader publicly denied a wall was planned — 59 days before building one.", "Primary"),
    ("S2", "Secret record", "1 Aug 1961", "USSR / GDR", "The real motive in private: skilled workers were fleeing and the economy was failing.", "Primary"),
    ("S3", "Newspaper", "14 Aug 1961", "East Germany", "The official public story: the Wall was a normal, defensive border control.", "Primary"),
    ("S4", "Propaganda", "1961", "East Germany", "How the state sold the Wall to its own people — as a victory over “spies”.", "Primary"),
    ("S5", "Military report", "17 Aug 1962", "East Germany", "The shoot-to-kill order in practice: 35 shots, an 18-year-old left to die.", "Primary"),
    ("S6", "Speech", "26 Jun 1963", "USA / West", "The Western argument: a state that must lock people in has already lost.", "Primary"),
    ("S7", "Statistics", "1949–90", "Neutral / research", "The Wall worked: emigration fell about 90% in one year, 1961 to 1962.", "Secondary"),
    ("S8", "Speech", "12 Jun 1987", "USA / West", "26 years on, the Wall was still a public weapon in the Cold War argument.", "Primary"),
    ("S9", "Government decree", "9 Nov 1989", "East Germany", "The Wall opened partly by accident — the decree was meant to start a day later.", "Primary"),
    ("S10", "Research database", "2017–19", "Neutral / research", "The human cost, established by evidence: at least 140 people killed at the Wall.", "Secondary"),
]
t = doc.add_table(rows=1, cols=6)
t.style = "Table Grid"
for i, h in enumerate(["#", "Source type", "Year", "Whose side", "What it proves", "P / S"]):
    c = t.rows[0].cells[i]
    c.text = ""
    run(c.paragraphs[0], h, bold=True, size=8.5, color=RGBColor(0xFF, 0xFF, 0xFF))
    shade(c, "1A1A1A")
for row in CHART:
    c = t.add_row().cells
    for i, v in enumerate(row):
        c[i].text = ""
        run(c[i].paragraphs[0], v, bold=(i == 0), size=8.5)

p = doc.add_paragraph()
rich(p, [("A good line to finish your chart with: ", "b"),
         ("the strongest thing about this set of sources is that S1 and S2 are ", ""),
         ("both East German", "i"),
         (" — and they contradict each other. What a government said in public (S1, S3, S4) and what it said in "
          "private (S2, S5, S9) are two different histories. That contrast is the argument your chart can make.", "")])

# ---------------- numbers ----------------
heading("The Wall by numbers")
NUMS = [("Total border around West Berlin", "155 km", "Height of the concrete wall", "3.6 m"),
        ("Running through the city itself", "43.1 km", "Length of concrete wall", "106 km"),
        ("Around West Berlin's outer edge", "111.9 km", "Anti-vehicle trenches", "105.5 km"),
        ("Watchtowers", "302", "Contact / signal fence", "127.5 km"),
        ("Bunkers", "20", "Patrol road (Kolonnenweg)", "124.3 km")]
t = doc.add_table(rows=0, cols=4)
t.style = "Table Grid"
for row in NUMS:
    c = t.add_row().cells
    for i, v in enumerate(row):
        c[i].text = ""
        run(c[i].paragraphs[0], v, bold=(i % 2 == 1), size=9.5)
p = doc.add_paragraph()
run(p, "Source: Berlin Wall Foundation (Stiftung Berliner Mauer) border-installation figures.", size=8.5, color=GREY)

# ---------------- bibliography ----------------
heading("Where every source came from")
p = doc.add_paragraph()
run(p, "Every one of these is a real, publicly accessible document. Type the address into a browser to see the original.",
    size=9, color=GREY)

BIB = [
    ("S1", "Ulbricht press conference, 15 June 1961",
     "germanhistorydocs.org/en/two-germanies-1961-1989/walter-ulbrichts-speech-on-the-berlin-question-june-15-1961\n"
     "berlingeschichte.de/bms/bmstxt01/0106dokb.htm"),
    ("S2", "Khrushchev–Ulbricht conversation, 1 Aug 1961",
     "digitalarchive.wilsoncenter.org/document/notes-conversation-comrade-ns-khrushchev-comrade-w-ulbricht-1-august-1961"),
    ("S3", "Neues Deutschland front page, 14 Aug 1961",
     "germanhistorydocs.org/de/zwei-deutsche-staaten-1961-1989/kommunistische-rechtfertigung-der-teilung-berlins-14-august-1961\n"
     "nd-archiv.de/ausgabe/1961-08-14"),
    ("S4", "SED pamphlet, “In Praise of the Berlin Wall”, 1961",
     "research.calvin.edu/german-propaganda-archive/schlugs13.htm"),
    ("S5", "Border brigade report on Peter Fechter, 17 Aug 1962",
     "chronik-der-mauer.de/todesopfer/171497/\nchronik-der-mauer.de/en/victims/180502/fechter-peter"),
    ("S6", "Kennedy, Rudolph Wilde Platz, 26 June 1963",
     "millercenter.org/the-presidency/presidential-speeches/june-26-1963-ich-bin-ein-berliner-speech\n"
     "jfklibrary.org/archives/other-resources/john-f-kennedy-speeches/berlin-w-germany-rudolph-wilde-platz-19630626"),
    ("S7", "East–West German migration statistics",
     "germanhistorydocs.org/en/two-germanies-1961-1989/east-west-german-immigration-statistics-1961-1990\n"
     "germanhistorydocs.org/en/occupation-and-the-emergence-of-two-states-1945-1961/refugee-movement-1950-1963"),
    ("S8", "Reagan, Brandenburg Gate, 12 June 1987",
     "reaganlibrary.gov/archives/speech/remarks-east-west-relations-brandenburg-gate-west-berlin\n"
     "archives.gov/publications/prologue/2007/summer/berlin.html"),
    ("S9", "GDR travel decree, 9 Nov 1989",
     "chronik-der-mauer.de/en/material/180400/press-conference-with-guenter-schabowski-9-november-1989"),
    ("S10", "Victims at the Berlin Wall, 1961–1989",
     "chronik-der-mauer.de/en/victims/\nstiftung-berliner-mauer.de/en/topics/victims-berlin-wall"),
]
t = doc.add_table(rows=1, cols=3)
t.style = "Table Grid"
for i, h in enumerate(["#", "Source", "Web address"]):
    c = t.rows[0].cells[i]
    c.text = ""
    run(c.paragraphs[0], h, bold=True, size=8.5, color=RGBColor(0xFF, 0xFF, 0xFF))
    shade(c, "1A1A1A")
for sid, name, urls in BIB:
    c = t.add_row().cells
    c[0].text = ""
    run(c[0].paragraphs[0], sid, bold=True, size=9)
    c[1].text = ""
    run(c[1].paragraphs[0], name, size=9)
    c[2].text = ""
    for j, u in enumerate(urls.split("\n")):
        p = c[2].paragraphs[0] if j == 0 else c[2].add_paragraph()
        run(p, u, size=8, color=RGBColor(0x1F, 0x4E, 0x79))

doc.add_paragraph()
p = doc.add_paragraph()
rich(p, [("Photographs. ", "b"),
         ("All eight photographs in photo-sheet.html come from Wikimedia Commons. They are either public domain (the CIA "
          "photographs, works of the US federal government) or CC BY-SA (the German Federal Archive / Bundesarchiv "
          "photographs, which must be credited). The photographer and licence are printed under each photo on the sheet.", "")],
     size=9)
p = doc.add_paragraph()
rich(p, [("Not included: ", "b"),
         ("the well-known photograph of Peter Fechter dying in the border strip is held by the Berlin police historical "
          "collection and is still in copyright, so it cannot be reproduced. You can view it on the Chronik der Mauer page "
          "listed under S5.", "")], size=9)

doc.save("berlin-wall-sources.docx")
print("saved berlin-wall-sources.docx")
